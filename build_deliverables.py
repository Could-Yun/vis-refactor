from pathlib import Path
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import zipfile,json
ROOT=Path(__file__).resolve().parent;out=ROOT/'交付物';out.mkdir(exist_ok=True)
doc=Document();sec=doc.sections[0];sec.top_margin=sec.bottom_margin=Cm(2.3);sec.left_margin=sec.right_margin=Cm(2.4)
# Remove template decorations and theme colours to keep the report black and white.
for st in doc.styles:
 for node in list(st.element.iter(qn('w:pBdr'))):node.getparent().remove(node)
 for node in st.element.iter(qn('w:color')):
  node.attrib.clear();node.set(qn('w:val'),'000000')
normal=doc.styles['Normal'];normal.font.name='Times New Roman';normal.font.size=Pt(12);normal.element.rPr.rFonts.set(qn('w:eastAsia'),'宋体');normal.paragraph_format.line_spacing=1.5
for h in ['Title','Heading 1','Heading 2']:
 st=doc.styles[h];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0);st.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'黑体');st.font.size=Pt(16 if h=='Title' else 14)
doc.add_heading('《可视化导论》课程网站重构\n关键亮点与修改说明',0)
def section(title,text):
 doc.add_heading(title,1)
 for s in text: p=doc.add_paragraph(s);p.paragraph_format.first_line_indent=Cm(.85)
section('一、项目定位与内容保留',['本项目针对原课程网站中资源难找、长篇内容定位困难、跨章节知识联系不清晰的问题，重建课程门户、章节导航与演示入口。保留五章教学正文、关键代码、既有数据集和讲义，并以交互工具帮助学生理解可视化设计的选择依据。','教学内容来源为 cinger007/vis，原站作者与授课教师信息保留在门户及页脚。本项目新增的功能属于教学呈现方式与交互组织的改进。'])
section('二、以知识地图建立课程总览',['首页的课程知识地图把数据、流程、色彩、低维可视化和时序可视化连接起来。节点对应概念，连线呈现概念之间的依赖；点击节点可进入相关章节，悬停用于查看关系。它将原本分散的章节入口转化为可探索的知识结构。'])
section('三、全站搜索与学习进度',['搜索覆盖章节与演示内容。中文采用二元组切分，英文采用词项匹配，按当前语言筛选结果，并提供命中片段与直达链接。用户可以通过搜索按钮和键盘打开搜索，使用方向键、回车和 Escape 完成选择与关闭。','各章使用分段式教学界面，因此学习进度按已访问小节记录，而非把滚动距离等同于掌握程度。首页进度环和继续阅读入口调用同一份浏览器本地记录；这些记录只代表访问情况，不作为学习成绩。'])
section('四、交互式通道排序与图表选型',['Mackinlay 通道排序工具根据定量、有序和名义三种数据类型切换排序，并提供解释。图形变化帮助学生理解：同一个视觉通道对于不同数据类型，表达能力并不相同。','新补充的图表选型工具从字段类型与分析任务出发，展示条形图、散点图、折线图、直方图与热力图的预览，解释位置、长度、明度等编码的适用条件，并给出返回相关课程的入口。工具接入首页和第 1、2、4 章，形成理论到设计决策的回路。'])
section('五、同数据色图对比实验',['色彩对比器在完全相同的 0—1 数值渐变上并排展示 jet、viridis 和 RdBu，并绘制相对亮度曲线。数值不变而色图改变，使学生能够区分数据边界与由不均匀亮度引入的视觉边界。','观察条件支持原色、灰度及三类色觉缺陷的近似模拟。页面明确说明模拟的教学性质。Viridis 用于连续数值，RdBu 需要有意义的中点；工具并未将某一种色图宣称为适合所有任务。'])
section('六、课程与 Demo 的可用性改进',['五章入口与 Demo 画廊统一组织，保持原课程路径。对于源仓库没有对应页面或原始体数据的两个演示，移除死链并保留原因说明，避免承诺不存在的功能。','回归示例原先只定义算法而没有实际图形，现补充散点图与拟合线。拖动样本点可重新计算最小二乘回归及 R²，使算法从代码片段变成可操作的教学演示。D3 与回归插件均保存在站点内，避免核心图表必须依赖外部 CDN。'])
section('七、双语、布局与可访问性',['正文保留原有中英文内容，工具的按钮、说明、图形标签随语言状态更新。章节间翻页标题根据目标章节统一校正。响应式布局使图表选型与色彩图组在窄屏堆叠，搜索框提供键盘交互。','图表同时呈现文字解释与图形预览。色觉模拟、灰度观察和亮度曲线用于提示颜色编码的局限，而不把颜色作为唯一的信息通道。'])
section('八、运行、验证与提交材料',['项目为静态 HTML、CSS 与 JavaScript，无需构建。解压源码后在站点目录执行 python -m http.server 8080，再访问 http://localhost:8080/。本地运行需要 HTTP 服务，不建议直接双击文件。','交付物目录保存验收报告、页面截图、源码压缩包、关键亮点说明及五分钟讲解稿。验收报告记录实际浏览器执行结果；讲解视频仍需提交者本人录制并配音。正式网址为 https://could-yun.github.io/vis-refactor/，以发布后的 HTTP 检查结果为准。'])
for filename,caption in [('色彩对比器.png','图 1　同数据的色彩与亮度对比工具')]:
 if (out/filename).exists():doc.add_picture(str(out/filename),width=Cm(16));p=doc.add_paragraph(caption);p.alignment=1
doc.save(out/'vis-refactor-关键亮点说明.docx')
script='''# 网页演示讲稿（约五分钟）

网址：https://could-yun.github.io/vis-refactor/
录制前检查网络、浏览器缩放与语言状态。括号内为操作提示，不必朗读。

## 0:00—0:35 项目目标
大家好，我展示的是《可视化导论》课程网站重构。改进围绕三个学习问题：怎样找到需要的内容，怎样定位当前学习位置，以及怎样建立章节之间的联系。五章教学正文、代码、数据和讲义继续保留，改动重点是内容组织和交互方式。

## 0:35—1:15 知识地图
（展示首页知识地图，悬停并点击一个概念。）首页从单纯章节入口扩展成知识总览。五章与核心概念构成网络，连线表示依赖，点击后进入相应章节。例如数据尺度会影响视觉通道，再进一步影响图表选择。学习者可以从问题出发，沿关系找到需要补充的知识。

## 1:15—1:55 搜索与继续阅读
（打开搜索，输入“色彩”，点击结果；切换小节后返回门户。）搜索把分散在章节和 Demo 中的内容统一索引，展示标题、位置和命中片段。中文支持连续词语匹配，英文也能搜索。章节访问记录保存在当前浏览器，首页用进度环展示已经访问的小节，并提供继续阅读。进度表示访问，而不是掌握程度。

## 1:55—2:45 通道排序与图表选型
（打开第二章通道排序，切换三种数据类型；展示首页选型工具。）定量、有序和名义数据的通道优先级不同。切换后图形与解释同步更新。选型工具把这一原则进一步变成操作：选字段类型，就能看到对应任务、推荐图表、编码理由和预览。例如类别加数值适合比较各组数量，两个数值适合用位置观察关联。推荐同时给出限制，避免把自动选图理解成唯一正确答案。

## 2:45—3:40 色彩对比
（第三章切换灰度与色觉条件。）这三幅图使用完全相同的数据，只改变色图。下方曲线表示相对亮度，因此可以直接观察色图是否制造了数据里不存在的边界。Viridis 面向连续数值，RdBu 用于有意义中点两侧的差异。灰度与色觉近似模拟帮助检查颜色是否承担了过多信息，也提醒我们不要只靠色相传达数值。

## 3:40—4:20 Demo 与代码学习
（进入 Demo 画廊，打开回归演示并拖动一个点。）演示集合现在有统一入口。回归示例补上了真实散点和拟合线，拖动点会更新拟合结果与 R 平方。学生能够把代码、数据与图形变化对应起来。原来缺少页面或数据的演示不再保留无效入口，页面说明了移除原因。

## 4:20—5:00 双语与收尾
（切到英文，再回中文，展示章节翻页与讲义。）正文、搜索和新增工具随语言状态切换；窄屏下图形堆叠，便于阅读。D3 与回归插件本地保存，减少外部服务对课堂使用的影响。这个项目的主要价值是让课程资源更容易被找到，让可视化理论能够通过操作被理解。谢谢。
'''
(out/'vis-refactor-五分钟演示讲稿.md').write_text(script,encoding='utf-8')
(ROOT/'README.md').write_text('# 可视化导论课程重构\n\n原课程：https://cinger007.github.io/vis/\n\n发布：https://could-yun.github.io/vis-refactor/\n\n运行：`python -m http.server 8080`，访问 http://localhost:8080/。\n\n功能：五章课程、双语讲义、全站搜索、学习进度、知识地图、通道排序、图表选型、色彩对比及 Demo 画廊。D3 7.9.0 与 d3-regression 1.3.10 已本地化。\n\n`交付物/` 含文档、演示讲稿和验收报告。`qa_site.py` 需要 Python Playwright 与 Edge，用于浏览器回归测试。\n',encoding='utf-8')
(ROOT/'.gitignore').write_text('交付物/\n__pycache__/\n*.log\n',encoding='utf-8')
with zipfile.ZipFile(out/'vis-refactor-网站源码.zip','w',zipfile.ZIP_DEFLATED) as z:
 for p in ROOT.rglob('*'):
  if p.is_file() and not any(x in p.relative_to(ROOT).parts for x in ['交付物','.git','__pycache__']):z.write(p,'vis-refactor/'+p.relative_to(ROOT).as_posix())
print('Highlights DOCX, script, README and source ZIP created.')
